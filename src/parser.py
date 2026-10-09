import re
from pathlib import Path
import hashlib

from pypdf import PdfReader


SUPPORTED = {".pdf", ".docx", ".txt"}


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return "\n".join(
            page.extract_text() or ""
            for page in PdfReader(str(path)).pages
        ).strip()

    if suffix == ".txt":
        return path.read_text(
            encoding="utf-8",
            errors="replace",
        ).strip()

    if suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        return "\n".join(parts).strip()

    raise ValueError(f"Unsupported file type: {suffix}")


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ").replace("\u200b", "")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def file_digest(path: Path) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(65536), b""):
            digest.update(chunk)

    return digest.hexdigest()

SKILLS = {
    "Python": r"\bpython\b",
    "FastAPI": r"\bfastapi\b",
    "Flask": r"\bflask\b",
    "Django": r"\bdjango\b",
    "PostgreSQL": r"\bpostgres(?:ql)?\b",
    "Redis": r"\bredis\b",
    "SQL": r"\bsql\b",

    "LangChain": r"\blangchain\b",
    "LangGraph": r"\blanggraph\b",
    "LlamaIndex": r"\bllamaindex\b",
    "RAG": r"\brag\b|retrieval[ -]augmented generation",
    "LLM": r"\bllms?\b|large language models?",
    "Embeddings": r"\bembeddings?\b",
    "Vector search": (
        r"vector (?:database|search|store)|"
        r"\bchroma(?:db)?\b|\bfaiss\b|\bpinecone\b|\bweaviate\b"
    ),
    "Agent/tool calling": (
        r"tool[ -]calling|\bagents?\b|function calling|"
        r"multi[ -]agent|google adk"
    ),

    "Docker": r"\bdocker\b",
    "GCP": r"\bgcp\b|google cloud",
    "AWS": r"\baws\b|amazon web services",
    "Azure": r"\bazure\b",
    "React": r"\breact(?:\.js)?\b",
    "Next.js": r"\bnext\.js\b",
    "Node.js": r"\bnode\.js\b|\bnodejs\b",

    "Testing": r"\bpytest\b|unit tests?|integration tests?",
    "Caching": r"\bcach(?:e|ing|ed)\b",
    "Async/concurrency": (
        r"\basync(?:io)?\b|concurren(?:cy|t)|"
        r"parallel processing|\bthreading\b"
    ),
}


# Common resume section headings.

HEAD = re.compile(
    r"(?im)^\s*(?:"
    r"projects?|technical projects?|selected projects?|"
    r"personal projects?|academic projects?|"
    r"work experience|professional experience|"
    r"relevant experience|research experience|"
    r"internships?|employment history|work history|"
    r"experience|technical skills|skills|"
    r"education|certifications|achievements|"
    r"summary|professional summary|about me"
    r")\s*:?\s*$"
)


def extract_candidate(text, filename):
    """
    Extract candidate details, skills, and project/work evidence
    from resume text.

    Important:
    If a project/work section cannot be identified, project evidence
    remains empty instead of falling back to the entire resume.
    This helps prevent skills-list keywords from qualifying a candidate.
    """

    if not isinstance(text, str):
        text = ""

    # Normalize lines while retaining the original text for evidence.
    lines = [
        re.sub(r"\s+", " ", line).strip(" |•\t")
        for line in text.splitlines()
        if line.strip()
    ]

    # Extract email addresses.
    emails = re.findall(
        r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b",
        text,
        re.IGNORECASE,
    )

    email_match = re.search(
        r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b",
        text,
        re.IGNORECASE,
    )

    # Try to identify the candidate's name near the beginning
    # of the resume or immediately before the email address.
    name = None

    if email_match:
        preceding_text = text[:email_match.start()]

        for line in preceding_text.splitlines()[:5]:
            line = re.sub(r"[|•—–:+]", " ", line).strip()
            words = re.findall(r"[A-Za-z][A-Za-z'.-]*", line)

            if (
                2 <= len(words) <= 4
                and not re.search(
                    r"developer|engineer|student|summary|education|"
                    r"skills|experience|linkedin|github|hyderabad|"
                    r"bengaluru|india|phone",
                    line,
                    re.IGNORECASE,
                )
            ):
                name = " ".join(words)
                break

    locations = {
        "bengaluru",
        "bangalore",
        "hyderabad",
        "ghaziabad",
        "india",
        "karnataka",
        "delhi",
        "mumbai",
        "pune",
        "chennai",
        "up",
        "telangana",
        "maharashtra",
    }

    if not name:
        for index, line in enumerate(lines[:12]):
            if (
                len(line) > 65
                or "@" in line
                or "," in line
                or re.search(
                    r"https?://|linkedin|github|phone|email",
                    line,
                    re.IGNORECASE,
                )
            ):
                continue

            words = re.findall(r"[A-Za-z][A-Za-z'.-]*", line)

            if (
                2 <= len(words) <= 4
                and not re.search(
                    r"developer|engineer|student|summary|education|"
                    r"skills|experience|profile|objective",
                    line,
                    re.IGNORECASE,
                )
            ):
                name = " ".join(words)
                break

            if (
                len(words) == 1
                and index + 1 < len(lines)
                and re.fullmatch(r"[A-Za-z'-]+", lines[index + 1])
            ):
                name = line + " " + lines[index + 1]
                break

    # Remove a trailing location if it was mistakenly included
    # as part of the detected name.
    if name:
        name_words = name.split()

        while (
            len(name_words) > 2
            and name_words[-1].lower().strip(".,") in locations
        ):
            name_words.pop()

        name = " ".join(name_words)

    # Fall back to the filename if no name could be extracted.
    if not name:
        name = filename.rsplit(".", 1)[0].replace("_", " ").title()

    # Extract GitHub username from a GitHub URL.
    github_usernames = re.findall(
        r"(?:https?://)?(?:www\.)?github\.com/([A-Za-z0-9-]+)",
        text,
        re.IGNORECASE,
    )

    github_usernames = [
        username
        for username in github_usernames
        if username.lower()
        not in {"features", "topics", "settings", "orgs", "marketplace"}
    ]

    github_url = (
        f"https://github.com/{github_usernames[0]}"
        if github_usernames
        else None
    )

    # Identify resume sections using recognized headings.
    matches = list(HEAD.finditer(text))
    sections = {}

    for index, match in enumerate(matches):
        start = match.end()

        end = (
            matches[index + 1].start()
            if index + 1 < len(matches)
            else len(text)
        )

        heading = match.group(0).strip().rstrip(":").lower()
        section_content = text[start:end].strip()

        # Keep repeated headings rather than discarding earlier sections.
        sections[heading] = (
            sections.get(heading, "") + "\n" + section_content
        ).strip()

    # Collect only project and work-related sections.
    project_sections = (
        "project",
        "experience",
        "internship",
    )

    project_evidence = "\n".join(
        content
        for heading, content in sections.items()
        if any(section in heading for section in project_sections)
    ).strip()

    # Do not fall back to the full resume when project evidence is
    # missing or too short. The eligibility checker can then reject
    # the resume for missing evidence, and it can be manually reviewed.

    # Identify skills across the full resume for ranking and reporting.
    matched_skills = [
        skill
        for skill, pattern in SKILLS.items()
        if re.search(pattern, text, re.IGNORECASE)
    ]

    return {
        "candidate_name": name,
        "email": emails[0] if emails else None,
        "github_url": github_url,
        "matched_skills": matched_skills,
        "text": text,
        "project_evidence_text": project_evidence,
        "source_sections": list(sections.keys()),
    }